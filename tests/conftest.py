from pathlib import Path

import pytest


@pytest.fixture
def documents(tmp_path: Path) -> list[Path]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.pdfgen.canvas import Canvas

    pdf = tmp_path / "report.pdf"
    canvas = Canvas(str(pdf))
    canvas.drawString(72, 720, "QuickMD fixture text")
    canvas.save()
    word = Document()
    word.add_paragraph("QuickMD fixture text")
    word.save(tmp_path / "report.docx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(1)
    ).text = "QuickMD fixture text"
    slides.save(tmp_path / "report.pptx")
    book = Workbook()
    book.active.append(["Title", "Value"])
    book.active.append(["QuickMD fixture text", 42])
    book.save(tmp_path / "report.xlsx")
    return [tmp_path / ("report" + ext) for ext in (".pdf", ".docx", ".pptx", ".xlsx")]
